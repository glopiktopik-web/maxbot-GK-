"""Сборка файлов-отчётов: Excel по посещаемости и Word по шаблонам."""

from __future__ import annotations

import logging
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

from config import DATA_DIR

logger = logging.getLogger(__name__)

EXPORT_DIR = DATA_DIR / "exports"

_SAFE = re.compile(r"[^\w\-.]+", re.UNICODE)


def _safe_name(text: str) -> str:
    return _SAFE.sub("_", text).strip("_") or "file"


def _prepare(name: str) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    return EXPORT_DIR / name


def attendance_xlsx(
    records: list[dict],
    absences: dict[int, list[str]],
    *,
    title: str,
    author: str | None = None,
) -> Path:
    """Сводный отчёт по посещаемости: журнал + сводка по студентам."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Журнал"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="44546A")
    wrap = Alignment(vertical="top", wrap_text=True)

    sheet.append(["Дата", "Группа", "Урок", "Дисциплина", "Отсутствовали", "Фамилии"])
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        names = absences.get(record["id"], [])
        sheet.append(
            [
                record["lesson_date"],
                record["group_name"],
                record["lesson_no"] or "",
                record["subject"] or "",
                record["absent_count"],
                ", ".join(names),
            ]
        )

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap

    widths = (12, 10, 7, 34, 15, 46)
    for number, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(number)].width = width
    sheet.freeze_panes = "A2"

    # ── сводка по студентам ───────────────────────────────────────────
    summary = workbook.create_sheet("Сводка")
    counter: Counter[tuple[str, str]] = Counter()
    for record in records:
        for student in absences.get(record["id"], []):
            counter[(record["group_name"], student)] += 1

    summary.append(["Группа", "Студент", "Пропущено занятий"])
    for cell in summary[1]:
        cell.font = header_font
        cell.fill = header_fill

    for (group, student), count in sorted(
        counter.items(), key=lambda item: (item[0][0], -item[1], item[0][1])
    ):
        summary.append([group, student, count])

    for number, width in enumerate((12, 34, 20), start=1):
        summary.column_dimensions[get_column_letter(number)].width = width
    summary.freeze_panes = "A2"

    # ── сводка по группам ─────────────────────────────────────────────
    totals = workbook.create_sheet("По группам")
    totals.append(["Группа", "Занятий отмечено", "Всего пропусков"])
    for cell in totals[1]:
        cell.font = header_font
        cell.fill = header_fill

    by_group: dict[str, list[int]] = {}
    for record in records:
        stats = by_group.setdefault(record["group_name"], [0, 0])
        stats[0] += 1
        stats[1] += record["absent_count"]

    for group, (lessons, absent) in sorted(by_group.items()):
        totals.append([group, lessons, absent])

    for number, width in enumerate((12, 20, 20), start=1):
        totals.column_dimensions[get_column_letter(number)].width = width

    workbook.properties.title = title
    if author:
        workbook.properties.creator = author

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    path = _prepare(f"Посещаемость_{_safe_name(title)}_{stamp}.xlsx")
    workbook.save(path)
    return path


def template_docx(title: str, body: str, *, author: str | None = None) -> Path:
    """Шаблон отчёта в виде документа Word, готового к заполнению."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(14)

    for block in body.split("\n"):
        text = block.rstrip()
        if not text:
            document.add_paragraph()
            continue
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(0)

    document.core_properties.title = title
    if author:
        document.core_properties.author = author

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    path = _prepare(f"{_safe_name(title)}_{stamp}.docx")
    document.save(path)
    return path
